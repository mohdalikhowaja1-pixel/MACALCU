import flet as ft
from datetime import datetime
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from storage.khata_db import init_khata_db, add_khata_entry, get_khata_entries, delete_khata_entry

def KhataScreen(on_back):
    try:
        init_khata_db()
    except Exception as e:
        print(f"DB Init Error: {e}")

    title_field = ft.TextField(label="Description / Item Name", border_radius=8, bgcolor="#0F172A", color="white")
    category_dropdown = ft.Dropdown(
        label="Category",
        options=[
            ft.dropdown.Option("Daily Sales"),
            ft.dropdown.Option("Cost of Goods (COGS)"),
            ft.dropdown.Option("Operational Expense"),
            ft.dropdown.Option("Other Revenue"),
        ],
        border_radius=8,
        bgcolor="#0F172A",
        color="white"
    )
    type_dropdown = ft.Dropdown(
        label="Entry Type",
        options=[
            ft.dropdown.Option("Sale (Revenue)"),
            ft.dropdown.Option("Expense (Cost)"),
        ],
        value="Sale (Revenue)",
        border_radius=8,
        bgcolor="#0F172A",
        color="white"
    )
    amount_field = ft.TextField(label="Amount (PKR)", keyboard_type=ft.KeyboardType.NUMBER, border_radius=8, bgcolor="#0F172A", color="white")
    date_field = ft.TextField(label="Date (YYYY-MM-DD)", value=datetime.now().strftime("%Y-%m-%d"), border_radius=8, bgcolor="#0F172A", color="white", expand=True)
    
    def on_filter_changed(e):
        try:
            refresh_ledger()
        except Exception as err:
            print(f"Filter error: {err}")

    report_filter = ft.Dropdown(
        label="Filter Report",
        options=[
            ft.dropdown.Option("All Time"),
            ft.dropdown.Option("This Month"),
        ],
        value="All Time",
        border_radius=8,
        bgcolor="#0F172A",
        color="white",
        on_select=on_filter_changed
    )

    status_text = ft.Text("", size=13)
    list_view = ft.Column(spacing=10)
    summary_card = ft.Container(bgcolor="#0F172A", padding=15, border_radius=12)

    def refresh_ledger():
        try:
            list_view.controls.clear()
            rows = get_khata_entries()
            
            total_sales = 0
            total_expenses = 0
            current_month_str = datetime.now().strftime("%Y-%m")

            for row in rows:
                r_id, title, category, entry_type, amount, date = row
                
                if report_filter.value == "This Month" and not date.startswith(current_month_str):
                    continue

                is_sale = "Sale" in entry_type
                if is_sale:
                    total_sales += amount
                    amt_color = "#34D399" 
                    prefix = "+"
                else:
                    total_expenses += amount
                    amt_color = "#F87171" 
                    prefix = "-"

                def make_delete(rid):
                    return lambda e: (delete_khata_entry(rid), refresh_ledger())

                list_view.controls.append(
                    ft.Container(
                        content=ft.Row([
                            ft.Column([
                                ft.Text(title, weight=ft.FontWeight.BOLD, color="white", size=15),
                                ft.Text(f"{category} • {date}", color="#94A3B8", size=12),
                            ], expand=True),
                            ft.Text(f"{prefix}PKR {amount:,.2f}", color=amt_color, weight=ft.FontWeight.BOLD, size=15),
                            ft.IconButton(icon=ft.Icons.DELETE_OUTLINE, icon_color="#EF4444", on_click=make_delete(r_id))
                        ], alignment=ft.MainAxisAlignment.SPACE_BETWEEN),
                        bgcolor="#1E293B",
                        padding=12,
                        border_radius=10
                    )
                )
            
            net_profit = total_sales - total_expenses
            profit_color = "#34D399" if net_profit >= 0 else "#F87171"
            summary_card.content = ft.Row([
                ft.Column([ft.Text("Total Revenue", color="#94A3B8", size=12), ft.Text(f"PKR {total_sales:,.2f}", color="#34D399", size=18, weight=ft.FontWeight.BOLD)]),
                ft.Column([ft.Text("Total Costs", color="#94A3B8", size=12), ft.Text(f"PKR {total_expenses:,.2f}", color="#F87171", size=18, weight=ft.FontWeight.BOLD)]),
                ft.Column([ft.Text("Net Profit", color="#94A3B8", size=12), ft.Text(f"PKR {net_profit:,.2f}", color=profit_color, size=18, weight=ft.FontWeight.BOLD)]),
            ], alignment=ft.MainAxisAlignment.SPACE_AROUND)
            
            if list_view.page:
                list_view.update()
            if summary_card.page:
                summary_card.update()
        except Exception as e:
            print(f"Error refreshing ledger: {e}")

    def handle_date_change(e):
        try:
            if e.control.value:
                date_field.value = e.control.value.strftime("%Y-%m-%d")
                date_field.update()
        except Exception as err:
            print(f"Date change error: {err}")

    date_picker = ft.DatePicker(on_change=handle_date_change)

    def open_date_picker(e):
        try:
            if date_picker not in e.page.overlay:
                e.page.overlay.append(date_picker)
            date_picker.open = True
            e.page.update()
        except Exception as err:
            status_text.value = "Could not open calendar picker."
            status_text.color = "#EF4444"
            status_text.update()

    # Professional Excel Report Exporter with Bold Headings & Fitted Columns
    def export_report_excel(e):
        try:
            rows = get_khata_entries()
            filename = f"Business_Khata_Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            filepath = os.path.join(os.path.expanduser("~"), "Downloads", filename)
            
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Business Khata"

            # Enable grid lines
            ws.views.sheetView[0].showGridLines = True

            # Header Styling
            headers = ["ID", "Description", "Category", "Type", "Amount (PKR)", "Date"]
            ws.append(headers)

            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
            header_alignment = Alignment(horizontal="center", vertical="center")
            thin_border = Border(
                left=Side(style='thin', color='CCCCCC'),
                right=Side(style='thin', color='CCCCCC'),
                top=Side(style='thin', color='CCCCCC'),
                bottom=Side(style='thin', color='CCCCCC')
            )

            for col_num in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_num)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
                cell.border = thin_border

            # Append Data Rows
            for r in rows:
                ws.append(list(r))

            # Format Data Rows
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(headers)):
                for cell in row:
                    cell.border = thin_border
                    cell.alignment = Alignment(vertical="center")
                    if cell.column == 5:  # Amount column formatting
                        cell.number_format = '#,##0.00'

            # Auto-fit Column Widths so nothing gets cut off or shows ###
            for col in ws.columns:
                max_len = 0
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                for cell in col:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                ws.column_dimensions[col_letter].width = max(max_len + 5, 12)

            wb.save(filepath)
            
            status_text.value = f"Formatted Excel report saved to Downloads!"
            status_text.color = "#34D399"
            status_text.update()
        except Exception as err:
            status_text.value = f"Export failed: {str(err)}"
            status_text.color = "#EF4444"
            status_text.update()

    def handle_add(e):
        try:
            if not title_field.value or not amount_field.value:
                status_text.value = "Please fill out description and amount."
                status_text.color = "#EF4444"
                status_text.update()
                return
            
            amt = float(amount_field.value)
            add_khata_entry(title_field.value, category_dropdown.value, type_dropdown.value, amt, date_field.value)
            title_field.value = ""
            amount_field.value = ""
            status_text.value = "Ledger entry added successfully!"
            status_text.color = "#34D399"
            status_text.update()
            refresh_ledger()
        except ValueError:
            status_text.value = "Invalid amount format! Please enter numbers only."
            status_text.color = "#EF4444"
            status_text.update()
        except Exception as err:
            status_text.value = f"An error occurred: {str(err)}"
            status_text.color = "#EF4444"
            status_text.update()

    refresh_ledger()

    return ft.ListView(
        expand=True,
        spacing=15,
        padding=10,
        controls=[
            ft.Row([
                ft.IconButton(icon=ft.Icons.ARROW_BACK, icon_color="white", on_click=on_back),
                ft.Text("Business Khata / Ledger", size=22, weight=ft.FontWeight.BOLD, color="white"),
                ft.Container(expand=True),
                report_filter
            ], alignment=ft.MainAxisAlignment.SPACE_BETWEEN),
            summary_card,
            ft.Container(
                content=ft.Column([
                    ft.Row([title_field, amount_field], spacing=10),
                    ft.Row([category_dropdown, type_dropdown], spacing=10),
                    ft.Row([
                        date_field,
                        ft.IconButton(
                            icon=ft.Icons.CALENDAR_MONTH,
                            icon_color="#38BDF8",
                            tooltip="Select Date",
                            on_click=open_date_picker
                        )
                    ], spacing=5, alignment=ft.MainAxisAlignment.SPACE_BETWEEN),
                    ft.Row([
                        ft.Button("Add Entry", icon=ft.Icons.ADD, on_click=handle_add, bgcolor="#6366F1", color="white"),
                        ft.Button("Export Excel Report", icon=ft.Icons.DOWNLOAD, on_click=export_report_excel, bgcolor="#10B981", color="white")
                    ], spacing=10),
                    status_text
                ], spacing=10),
                bgcolor="#1E293B", padding=16, border_radius=12
            ),
            ft.Divider(color="#334155"),
            ft.Text("Transaction History", size=16, weight=ft.FontWeight.BOLD, color="white"),
            list_view
        ]
    )