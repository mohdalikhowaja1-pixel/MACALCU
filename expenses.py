import os
import flet as ft
from datetime import datetime
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors

from storage.expense_storage import add_expense, get_all_expenses, delete_expense

def build_tracker_screen(on_navigate, currency="PKR"):
    placeholder_style = ft.TextStyle(size=12, color=ft.Colors.GREY_500)

    # Input Control Fields
    title_input = ft.TextField(
        label="Expense Description",
        hint_text="e.g. Monthly Grocery",
        hint_style=placeholder_style,
        width=350,
    )

    amount_input = ft.TextField(
        label="Amount",
        hint_text="e.g. 2500",
        hint_style=placeholder_style,
        keyboard_type=ft.KeyboardType.NUMBER,
        prefix=ft.Text(f"{currency} "),
        width=350,
    )

    categories = [
        "Food & Dining",
        "Groceries",
        "Rent & Mortgage",
        "Bills & Utilities",
        "Transportation & Fuel",
        "Shopping & Clothing",
        "Healthcare & Medical",
        "Entertainment & Leisure",
        "Education & Books",
        "Travel & Vacation",
        "Investments & Savings",
        "Personal Care",
        "Other",
    ]

    category_dropdown = ft.Dropdown(
        label="Category",
        options=[ft.dropdown.Option(c) for c in categories],
        value="Food & Dining",
        width=350,
    )

    date_input = ft.TextField(
        label="Date (YYYY-MM-DD)",
        hint_text=f"e.g. {datetime.now().strftime('%Y-%m-%d')}",
        hint_style=placeholder_style,
        width=290,
    )

    filter_date_input = ft.TextField(
        label="Filter List by Date (YYYY-MM-DD)",
        hint_text="e.g. 2026-09-29 or empty for All",
        hint_style=placeholder_style,
        width=240,
    )

    # Date Picker for Adding Expenses
    def on_add_date_change(e):
        if add_date_picker.value:
            date_input.value = add_date_picker.value.strftime("%Y-%m-%d")
            if date_input.page:
                date_input.page.update()

    add_date_picker = ft.DatePicker(
        first_date=datetime(2020, 1, 1),
        last_date=datetime(2030, 12, 31),
        on_change=on_add_date_change,
    )

    def open_add_date_picker(e):
        try:
            if hasattr(e.page, "open"):
                e.page.open(add_date_picker)
            else:
                if add_date_picker not in e.page.overlay:
                    e.page.overlay.append(add_date_picker)
                add_date_picker.open = True
                e.page.update()
        except Exception:
            pass

    # Date Picker for Filtering Expenses
    def on_filter_date_change(e):
        if filter_date_picker.value:
            filter_date_input.value = filter_date_picker.value.strftime("%Y-%m-%d")
            populate_expenses_list()
            if filter_date_input.page:
                filter_date_input.page.update()

    filter_date_picker = ft.DatePicker(
        first_date=datetime(2020, 1, 1),
        last_date=datetime(2030, 12, 31),
        on_change=on_filter_date_change,
    )

    def open_filter_date_picker(e):
        try:
            if hasattr(e.page, "open"):
                e.page.open(filter_date_picker)
            else:
                if filter_date_picker not in e.page.overlay:
                    e.page.overlay.append(filter_date_picker)
                filter_date_picker.open = True
                e.page.update()
        except Exception:
            pass

    error_text = ft.Text("", color=ft.Colors.RED_400, size=13)
    status_text = ft.Text("", color=ft.Colors.GREEN_400, size=13)
    total_summary = ft.Text(f"Total Spent: {currency} 0.00", size=18, weight=ft.FontWeight.BOLD, color=ft.Colors.BLUE_400)
    expenses_list_column = ft.Column(spacing=8)

    def populate_expenses_list():
        expenses_list_column.controls.clear()
        selected_date = filter_date_input.value.strip() if filter_date_input.value else None
        
        # Date Filter Safeguard
        if selected_date:
            try:
                datetime.strptime(selected_date, "%Y-%m-%d")
            except ValueError:
                error_text.value = "Filter date must be in YYYY-MM-DD format."
                records = get_all_expenses()
                selected_date = None
        
        records = get_all_expenses(filter_date=selected_date)
        total = 0.0

        for record in records:
            try:
                e_id, title, amount, category, date_str = record
                total += amount

                def make_delete_handler(id_to_del):
                    return lambda e: handle_delete(id_to_del, e)

                item_card = ft.Card(
                    content=ft.Container(
                        padding=10,
                        content=ft.Row([
                            ft.Column([
                                ft.Text(title, size=15, weight=ft.FontWeight.BOLD),
                                ft.Text(f"{category} • {date_str}", size=12, color=ft.Colors.GREY_400),
                            ], expand=True),
                            ft.Text(f"{currency} {amount:,.2f}", size=15, weight=ft.FontWeight.BOLD, color=ft.Colors.RED_400),
                            ft.IconButton(
                                icon=ft.Icons.DELETE_OUTLINED,
                                icon_color=ft.Colors.GREY_500,
                                tooltip="Delete Expense",
                                on_click=make_delete_handler(e_id)
                            )
                        ], alignment=ft.MainAxisAlignment.SPACE_BETWEEN)
                    )
                )
                expenses_list_column.controls.append(item_card)
            except Exception:
                continue

        total_summary.value = f"Total Spent: {currency} {total:,.2f}"

    def handle_add(e):
        error_text.value = ""
        status_text.value = ""
        
        # Field Safeguards
        if not title_input.value or not title_input.value.strip():
            error_text.value = "Please enter a valid expense description."
            e.page.update()
            return

        try:
            amt = float(amount_input.value.strip().replace(",", ""))
            if amt <= 0:
                raise ValueError()
        except (ValueError, AttributeError, TypeError):
            error_text.value = "Please enter a valid positive numerical amount."
            e.page.update()
            return

        entered_date = date_input.value.strip() if date_input.value else ""
        if entered_date:
            try:
                datetime.strptime(entered_date, "%Y-%m-%d")
                final_date = entered_date
            except ValueError:
                error_text.value = "Invalid date format! Use YYYY-MM-DD or pick from calendar."
                e.page.update()
                return
        else:
            final_date = datetime.now().strftime("%Y-%m-%d")

        try:
            selected_cat = category_dropdown.value if category_dropdown.value else "Other"
            add_expense(title_input.value.strip(), amt, selected_cat, final_date)

            title_input.value = ""
            amount_input.value = ""
            date_input.value = ""
            status_text.value = f"Expense added successfully for {final_date}!"
            populate_expenses_list()
        except Exception as ex:
            error_text.value = f"Failed to save expense: {str(ex)}"

        e.page.update()

    def handle_delete(expense_id, e):
        try:
            delete_expense(expense_id)
            populate_expenses_list()
        except Exception as ex:
            error_text.value = f"Error deleting entry: {str(ex)}"
        e.page.update()

    def handle_filter(e):
        error_text.value = ""
        populate_expenses_list()
        e.page.update()

    def clear_filter(e):
        error_text.value = ""
        filter_date_input.value = ""
        populate_expenses_list()
        e.page.update()

    def export_to_excel(e):
        try:
            records = get_all_expenses(filter_date=filter_date_input.value.strip() if filter_date_input.value else None)
            if not records:
                status_text.value = "No expenses to export!"
                e.page.update()
                return

            df = pd.DataFrame(records, columns=["ID", "Title", "Amount", "Category", "Date"])
            df = df[["Date", "Title", "Category", "Amount"]]
            
            filename = f"expenses_export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            
            with pd.ExcelWriter(filename, engine='openpyxl') as writer:
                df.to_excel(writer, index=False, sheet_name="Expense Report", startrow=1)
                worksheet = writer.sheets["Expense Report"]

                header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
                header_fill = PatternFill(start_color="2B2D42", end_color="2B2D42", fill_type="solid")
                bold_font = Font(name="Calibri", size=11, bold=True)
                thin_border = Border(
                    left=Side(style='thin', color='D3D3D3'),
                    right=Side(style='thin', color='D3D3D3'),
                    top=Side(style='thin', color='D3D3D3'),
                    bottom=Side(style='thin', color='D3D3D3')
                )

                worksheet["A1"] = "Expense Report"
                worksheet["A1"].font = Font(name="Calibri", size=14, bold=True)

                for col_num in range(1, 5):
                    cell = worksheet.cell(row=2, column=col_num)
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                total_sum = 0
                start_row = 3
                for row_idx, row in df.iterrows():
                    current_row = start_row + row_idx
                    total_sum += row["Amount"]
                    for col_idx in range(1, 5):
                        cell = worksheet.cell(row=current_row, column=col_idx)
                        cell.border = thin_border
                        if col_idx == 4:
                            cell.number_format = '#,##0.00'
                            cell.alignment = Alignment(horizontal="right")

                total_row = start_row + len(df)
                worksheet.cell(row=total_row, column=3, value="Total Spent:").font = bold_font
                worksheet.cell(row=total_row, column=3).alignment = Alignment(horizontal="right")
                total_cell = worksheet.cell(row=total_row, column=4, value=total_sum)
                total_cell.font = bold_font
                total_cell.number_format = '#,##0.00'
                total_cell.alignment = Alignment(horizontal="right")

                for col in worksheet.columns:
                    max_len = max(len(str(cell.value or '')) for cell in col)
                    col_letter = col[0].column_letter
                    worksheet.column_dimensions[col_letter].width = max(max_len + 5, 12)

            status_text.value = f"Exported to Excel: {os.path.abspath(filename)}"
        except Exception as ex:
            error_text.value = f"Export Excel failed: {str(ex)}"
        e.page.update()

    def export_to_pdf(e):
        try:
            records = get_all_expenses(filter_date=filter_date_input.value.strip() if filter_date_input.value else None)
            if not records:
                status_text.value = "No expenses to export!"
                e.page.update()
                return

            filename = f"expenses_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
            doc = SimpleDocTemplate(filename, pagesize=letter)
            elements = []

            styles = getSampleStyleSheet()
            title_style = styles['Heading1']
            elements.append(Paragraph("Expense Report", title_style))
            elements.append(Spacer(1, 12))

            table_data = [["Date", "Description", "Category", f"Amount ({currency})"]]
            total_sum = 0
            for r in records:
                table_data.append([r[4], r[1], r[3], f"{r[2]:,.2f}"])
                total_sum += r[2]

            table_data.append(["", "", "Total Spent:", f"{total_sum:,.2f}"])

            t = Table(table_data, colWidths=[90, 180, 130, 100])
            t.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#2B2D42")),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('ALIGN', (3, 0), (3, -1), 'RIGHT'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 10),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
                ('FONTNAME', (2, -1), (3, -1), 'Helvetica-Bold'),
            ]))

            elements.append(t)
            doc.build(elements)

            status_text.value = f"Exported PDF Report: {os.path.abspath(filename)}"
        except Exception as ex:
            error_text.value = f"Export PDF failed: {str(ex)}"
        e.page.update()

    populate_expenses_list()

    return ft.ListView(
        expand=True,
        spacing=10,
        padding=10,
        controls=[
            ft.Row([
                ft.IconButton(
                    icon=ft.Icons.ARROW_BACK,
                    on_click=lambda _: on_navigate("home")
                ),
                ft.Text("Expense Tracker", size=22, weight=ft.FontWeight.BOLD),
            ]),
            ft.Divider(),
            title_input,
            amount_input,
            category_dropdown,
            ft.Row([
                date_input,
                ft.IconButton(
                    icon=ft.Icons.CALENDAR_MONTH,
                    tooltip="Select Date from Calendar",
                    on_click=open_add_date_picker,
                ),
            ], spacing=5),
            error_text,
            status_text,
            ft.Button(content=ft.Text("Add Expense"), on_click=handle_add),
            ft.Divider(),
            ft.Row([
                filter_date_input,
                ft.IconButton(
                    icon=ft.Icons.CALENDAR_MONTH,
                    tooltip="Select Filter Date from Calendar",
                    on_click=open_filter_date_picker,
                ),
                ft.Button(content=ft.Text("Filter Date"), on_click=handle_filter),
                ft.Button(content=ft.Text("Show All"), on_click=clear_filter),
            ], alignment=ft.MainAxisAlignment.START, spacing=5),
            ft.Row([
                ft.Button(content=ft.Text("Export to Excel"), on_click=export_to_excel),
                ft.Button(content=ft.Text("Export to PDF"), on_click=export_to_pdf),
            ], spacing=10),
            ft.Divider(),
            total_summary,
            ft.Container(height=5),
            expenses_list_column,
        ]
    )